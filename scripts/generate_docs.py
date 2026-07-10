#!/usr/bin/env python3
"""Generate project documents for the email-system project.
Requires: python-docx (nix-shell -p python313Packages.python-docx)
"""

from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from datetime import datetime, timedelta
import os

PROJECT_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOCS_DIR = os.path.join(PROJECT_DIR, "docs")

def add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    return h

def add_para(doc, text, bold=False, size=None):
    p = doc.add_paragraph()
    run = p.add_run(text)
    if bold:
        run.bold = True
    if size:
        run.font.size = Pt(size)
    return p

def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers), style='Table Grid')
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    # Header
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            for run in p.runs:
                run.bold = True
    # Data
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            table.rows[ri + 1].cells[ci].text = str(val)
    doc.add_paragraph()
    return table


# ============================================================
# 1. Weekly Reports (weeks 2-5)
# ============================================================
def generate_weekly_report(week_num, start_date, end_date, summary, suggestions, next_plan):
    doc = Document()

    # Title
    title = doc.add_heading('工作周报', level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_para(doc, 'Work Review', size=10)

    # Info table
    add_table(doc, ['项目组', '姓名', '周次'],
              [['Email System 开发组', 'x12w', f'第{week_num}周 ({start_date} ~ {end_date})']])

    add_heading(doc, '本周综述', level=2)
    for item in summary:
        doc.add_paragraph(item, style='List Bullet')

    add_heading(doc, '本周建议和意见', level=2)
    for item in suggestions:
        doc.add_paragraph(item, style='List Bullet')

    add_heading(doc, '下周计划', level=2)
    for item in next_plan:
        doc.add_paragraph(item, style='List Bullet')

    path = os.path.join(DOCS_DIR, f'3.个人工作周报-第{week_num}周.docx')
    doc.save(path)
    print(f"  ✓ 第{week_num}周周报 → {path}")
    return path


def generate_all_weekly_reports():
    print("\n📝 生成个人工作周报...")

    # Week 2: June 9 - June 15
    generate_weekly_report(
        week_num=2,
        start_date='2026.06.09',
        end_date='2026.06.15',
        summary=[
            '完成前端核心页面开发：登录注册页面(LoginView.vue)和邮件主页(MailHomeView.vue)，集成Element Plus组件库，实现登录表单提交和注册功能。',
            '完成邮件业务后端三层架构：MailMessage、MailAccount、MailFolder等Entity实体类，Mapper数据访问层，以及Service业务逻辑层的搭建。',
            '实现JWT Token认证机制：TokenService签名/验签、JwtAuthenticationFilter拦截器、SecurityConfig安全配置，完成登录/注册/Token刷新API。',
            '开发邮件核心功能：邮件列表查询(分页/搜索/筛选)、邮件详情查看、发送邮件(SMTP)、保存草稿、软删除邮件。',
            '实现邮箱账号管理：添加/编辑/删除邮箱账号、SMTP/IMAP配置、账号归属校验。',
            '完成前端路由守卫和Pinia状态管理，实现登录态持久化和未登录拦截。',
        ],
        suggestions=[
            '建议引入前端错误边界处理，避免API异常导致页面白屏。',
            '密码存储需加密，当前authPasswordEncrypted字段实际存储明文，存在安全隐患。',
        ],
        next_plan=[
            '计划实现IMAP邮件同步功能，支持从第三方邮箱拉取邮件。',
            '计划开发附件上传/下载功能，集成MinIO对象存储。',
            '计划完善前端邮件列表交互（标记已读、星标、删除等操作）。',
        ]
    )

    # Week 3: June 16 - June 22
    generate_weekly_report(
        week_num=3,
        start_date='2026.06.16',
        end_date='2026.06.22',
        summary=[
            '实现IMAP邮件同步服务(MailSyncService)：支持从IMAP服务器拉取INBOX邮件，按Message-ID去重，首次同步拉取配置数量(默认50封)，后续增量拉取20封。',
            '完成附件功能：AttachmentController上传/下载/删除API，AttachmentService处理本地存储和MinIO对象存储两种模式，支持文件流下载。',
            '实现联系人管理模块：ContactController增删改查API，Contact实体含姓名/邮箱/电话/备注字段。',
            '开发文件夹管理：系统自动创建收件箱/发件箱/草稿箱/已删除文件夹，FolderController提供文件夹列表查询。',
            '完善前端API层：封装axios请求拦截器(自动附加Token)、响应拦截器(统一错误处理/401跳转登录)，完成auth/mail API模块。',
            '实现前端邮件发送界面：收件人/抄送/密送输入、富文本编辑器、附件上传、草稿保存。',
        ],
        suggestions=[
            'IMAP同步目前硬编码userId=1，需要改为遍历所有用户。',
            '发送邮件后的已发送副本应自动同步到IMAP Sent文件夹，保持多端一致。',
        ],
        next_plan=[
            '计划开发智能邮件分析插件系统，实现垃圾邮件检测和优先级分析。',
            '计划添加邮件推送事件机制，支持安全告警和高优先级通知。',
            '计划编写deploy.sh一键部署脚本，简化生产环境部署流程。',
        ]
    )

    # Week 4: June 23 - June 29
    generate_weekly_report(
        week_num=4,
        start_date='2026.06.23',
        end_date='2026.06.29',
        summary=[
            '实现智能分析插件系统：IntelligencePluginClient接口定义标准化I/O契约，RuleBasedIntelligencePluginClient基于规则引擎实现垃圾邮件检测、关键词匹配、钓鱼链接识别。',
            '开发智能分析服务(IntelligenceAnalysisServiceImpl)：集成插件调用、结果缓存、威胁指标提取，支持按消息ID分析和结果查询。',
            '实现邮件推送事件(PushEvent)：智能分析结果自动生成推送事件，支持安全告警/高优先级通知两种类型，前端可查看和管理推送事件。',
            '编写Python智能分析插件(plugins/intelligence/python)：实现analyzer.py分析引擎和plugin_entry.py入口，通过HTTP与Java后端通信。',
            '完成deploy.sh一键部署脚本：包含check/build/start/stop/restart/status/logs/update/env/dev等12个子命令。',
            '实现Nginx反向代理动态配置：自动检测SSL证书(/etc/letsencrypt)，存在时自动启用HTTPS，支持前端SPA路由和API代理。',
            '完成docker-compose.yml编排：MySQL 8.4 + Redis 7.4 + MinIO + Mailpit + docker-mailserver，含健康检查和数据持久化。',
        ],
        suggestions=[
            'deploy.sh中Nginx使用--network host存在安全隐患，应改为端口映射方式。',
            '智能分析结果Map缓存无上限，需要添加淘汰机制防止内存溢出。',
        ],
        next_plan=[
            '计划集成docker-mailserver实现独立IMAP/SMTP服务，支持多域名邮箱。',
            '计划完善错误处理和日志记录，添加全局异常拦截器。',
            '计划进行安全性加固：加密存储邮箱密码、修复CORS配置、添加Token黑名单等。',
            '计划编写项目文档：软件需求规约、架构文档、用例规约、项目总结报告。',
        ]
    )

    # Week 5: June 30 - July 6
    generate_weekly_report(
        week_num=5,
        start_date='2026.06.30',
        end_date='2026.07.06',
        summary=[
            '完成安全性全面加固(fix/security-hardening分支)：修复CORS配置漏洞(允许具体域名替代通配符)、实现AES-256-GCM加密存储IMAP/SMTP密码、添加Redis Token黑名单实现真正登出。',
            '修复多项关键Bug：MailSyncService.getFrom()空指针异常、IMAP同步线程池泄漏(new Thread改为ExecutorService)、RateLimitFilter内存泄漏(添加定时清理)、智能分析缓存无限增长。',
            '优化MailSyncService：硬编码userId=1改为遍历所有用户、时区从ZoneId.systemDefault()改为Asia/Shanghai显式指定、HTML预览剥离script/style块防止XSS。',
            '添加密码强度校验：自定义@ValidPassword注解，要求最少8位且包含字母和数字。',
            '修复deploy.sh部署脚本问题：SSL证书路径变量转义错误、Nginx改用--add-host替代--network host、proxy_pass使用host.docker.internal。',
            '完善docker-compose配置：添加IMAP_FETCH_COUNT环境变量、CORS允许域名配置。',
            '集成docker-mailserver(docker-mailserver:15.0.0)：配置SSL_TYPE=manual、ENABLE_IMAP=1、独立邮件服务器支持。',
            '完成答辩PPT生成脚本(generate_ppt.py)：11页精简版，含真实项目数据和部署信息，RGB颜色值溢出修复。',
        ],
        suggestions=[
            '建议引入Spring Boot Actuator进行应用健康监控和指标收集。',
            '建议添加API接口文档(Swagger/OpenAPI)，方便前后端协作。',
            '前端Token存储建议从localStorage迁移到HttpOnly Cookie，进一步增强安全性。',
        ],
        next_plan=[
            '计划实现邮件全文搜索功能(Elasticsearch集成)。',
            '计划开发邮件规则引擎：自动归档、标签分类、自动回复。',
            '计划添加多用户管理后台和权限控制(RBAC)。',
            '计划编写单元测试和集成测试，提升代码覆盖率。',
        ]
    )

    print("  ✅ 所有周报生成完毕")


# ============================================================
# 2. Software Requirements Specification
# ============================================================
def generate_srs():
    print("\n📝 生成软件需求规约...")
    doc = Document()

    # Title page
    title = doc.add_heading('Email System 邮件管理系统', level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_para(doc, '软件需求规约', size=14)
    add_para(doc, '版本 1.0', size=10)
    add_para(doc, f'日期: 2026-07-07', size=10)
    doc.add_page_break()

    # Revision history
    add_heading(doc, '修订历史记录', level=1)
    add_table(doc, ['日期', '版本', '说明', '作者'],
              [['2026/07/07', '1.0', '初始版本，完整需求规约', 'x12w']])

    # TOC placeholder
    add_heading(doc, '目录', level=1)
    add_para(doc, '1. 简介\n2. 整体说明\n3. 具体需求\n4. 用户界面原型')

    # 1. Introduction
    add_heading(doc, '1. 简介', level=1)
    add_heading(doc, '1.1 目的', level=2)
    add_para(doc, '本文档旨在详细描述Email System邮件管理系统的软件需求，为系统的设计、开发和测试提供依据。该系统是一个基于Web的全栈邮件管理平台，支持多邮箱账号管理、IMAP邮件同步、SMTP邮件发送、智能邮件分析等功能。')

    add_heading(doc, '1.2 范围', level=2)
    add_para(doc, 'Email System是一个企业级邮件管理系统，涵盖以下核心功能模块：\n'
               '• 用户认证与授权：注册、登录、JWT Token管理\n'
               '• 邮箱账号管理：多账号SMTP/IMAP配置与管理\n'
               '• 邮件收发：SMTP发送、IMAP同步接收、附件上传下载\n'
               '• 文件夹管理：收件箱/发件箱/草稿箱/已删除等系统文件夹\n'
               '• 联系人管理：通讯录增删改查\n'
               '• 智能分析：垃圾邮件检测、优先级分析、威胁识别\n'
               '• 系统管理：一键部署、健康检查、日志管理')

    add_heading(doc, '1.3 定义、首字母缩写词和缩略语', level=2)
    add_table(doc, ['术语/缩略语', '说明'],
              [['SMTP', 'Simple Mail Transfer Protocol，简单邮件传输协议，用于发送邮件'],
               ['IMAP', 'Internet Message Access Protocol，互联网消息访问协议，用于接收邮件'],
               ['JWT', 'JSON Web Token，用于无状态身份认证'],
               ['MinIO', '开源对象存储服务，用于附件存储'],
               ['Mailpit', '邮件测试工具，用于开发环境SMTP/IMAP模拟'],
               ['SRS', 'Software Requirements Specification，软件需求规约']])

    add_heading(doc, '1.4 参考资料', level=2)
    add_para(doc, '1. Spring Boot 3.4 官方文档\n'
               '2. Vue 3 + Element Plus 官方文档\n'
               '3. Jakarta Mail API 规范\n'
               '4. MyBatis-Plus 3.5 文档\n'
               '5. Docker 官方文档')

    add_heading(doc, '1.5 概述', level=2)
    add_para(doc, '本文档后续章节包括：整体说明（用例模型、假设与依赖）、具体需求（功能需求、非功能需求、接口需求）、用户界面原型。')

    # 2. Overall Description
    add_heading(doc, '2. 整体说明', level=1)
    add_heading(doc, '2.1 用例模型调查', level=2)
    add_table(doc, ['主要参与者', '优先级', '用例名', '用例概述'],
              [['注册用户', '1', '用户注册与登录', '用户注册账号、登录系统、Token刷新与登出'],
               ['注册用户', '1', '邮箱账号管理', '添加/编辑/删除邮箱账号，配置SMTP/IMAP参数'],
               ['注册用户', '1', '邮件收发', '发送邮件、查看收件箱、查看邮件详情、删除邮件'],
               ['注册用户', '1', 'IMAP同步', '自动从IMAP服务器同步邮件到本地数据库'],
               ['注册用户', '2', '联系人管理', '管理通讯录联系人信息'],
               ['注册用户', '2', '附件管理', '上传附件、下载附件、删除附件'],
               ['注册用户', '2', '智能分析', '查看邮件垃圾检测、优先级分析、威胁识别结果'],
               ['系统管理员', '2', '系统部署', '一键部署、服务启停、状态监控、日志查看']])

    add_heading(doc, '2.2 假设与依赖关系', level=2)
    add_para(doc, '假设条件：\n'
               '• 用户拥有可正常使用的邮箱账号（SMTP/IMAP服务可用）\n'
               '• 服务器具备Java 17+运行环境和Docker环境\n'
               '• 网络环境支持外部SMTP/IMAP服务的访问\n\n'
               '依赖关系：\n'
               '• 系统依赖MySQL 8.4数据库存储结构化数据\n'
               '• 系统依赖Redis 7.4进行缓存和Token黑名单管理\n'
               '• 系统依赖MinIO进行附件对象存储\n'
               '• 邮件发送依赖外部SMTP服务器或内置docker-mailserver')

    # 3. Specific Requirements
    add_heading(doc, '3. 具体需求', level=1)
    add_heading(doc, '3.1 功能需求', level=2)

    # FR list
    fr_data = [
        ['FR-001', '用户注册', '支持用户名+密码+邮箱注册，密码至少8位含字母数字', '高'],
        ['FR-002', '用户登录', '用户名+密码登录，返回JWT AccessToken和RefreshToken', '高'],
        ['FR-003', 'Token刷新', '使用RefreshToken刷新AccessToken，无需重新登录', '高'],
        ['FR-004', '安全登出', '登出时Token加入Redis黑名单，立即失效', '高'],
        ['FR-005', '邮箱账号CRUD', '添加/编辑/删除邮箱账号，配置SMTP/IMAP参数', '高'],
        ['FR-006', '发送邮件', '通过SMTP发送邮件，支持To/CC/BCC、HTML内容、附件', '高'],
        ['FR-007', 'IMAP同步', '定时从IMAP服务器同步INBOX邮件，支持Message-ID去重', '高'],
        ['FR-008', '邮件列表', '分页查询邮件列表，支持文件夹筛选、关键词搜索、已读过滤', '高'],
        ['FR-009', '邮件详情', '查看邮件完整内容（发件人、收件人、主题、正文、附件）', '高'],
        ['FR-010', '邮件操作', '标记已读/未读、软删除（移入回收站）', '中'],
        ['FR-011', '草稿保存', '未发送邮件保存为草稿，可继续编辑和发送', '中'],
        ['FR-012', '附件管理', '上传附件到MinIO/本地存储，支持下载和删除', '中'],
        ['FR-013', '联系人管理', '通讯录增删改查，支持姓名/邮箱/电话/备注', '中'],
        ['FR-014', '文件夹管理', '自动创建系统文件夹，显示各文件夹邮件数量', '中'],
        ['FR-015', '智能分析', '分析邮件垃圾概率、优先级、安全风险，生成威胁指标', '低'],
        ['FR-016', '推送事件', '智能分析结果生成推送通知，支持已读标记', '低'],
        ['FR-017', '速率限制', '登录/注册接口IP级别速率限制，防止暴力破解', '高'],
        ['FR-018', '系统部署', '一键部署脚本：check/build/start/stop/restart/status/logs', '中'],
    ]
    add_table(doc, ['编号', '功能名称', '描述', '优先级'], fr_data)

    add_heading(doc, '3.2 非功能需求', level=2)
    add_table(doc, ['类别', '需求描述'],
              [['性能', '邮件列表查询响应时间 < 500ms (P95)，支持100+并发用户'],
               ['安全', 'JWT Token认证、密码BCrypt加密、邮箱密码AES-256-GCM加密存储、IP速率限制'],
               ['可用性', '系统支持Docker Compose一键部署，基础设施自动健康检查'],
               ['可扩展性', '智能分析采用插件架构，支持Python/原生插件扩展'],
               ['兼容性', '支持主流浏览器(Chrome/Firefox/Edge)，后端Java 17+，前端Node 18+'],
               ['数据持久性', 'MySQL数据卷持久化、MinIO对象存储、Redis RDB持久化']])

    add_heading(doc, '3.3 接口需求', level=2)
    add_para(doc, '系统提供RESTful API接口，前缀为 /api/，主要接口如下：\n'
               '• /api/auth/* — 认证接口（登录、注册、刷新、登出、获取当前用户）\n'
               '• /api/mail-accounts/* — 邮箱账号管理\n'
               '• /api/messages/* — 邮件收发与管理\n'
               '• /api/folders/* — 文件夹管理\n'
               '• /api/contacts/* — 联系人管理\n'
               '• /api/attachments/* — 附件管理\n'
               '• /api/intelligence/* — 智能分析\n'
               '• /api/health — 健康检查\n\n'
               '所有接口（除登录/注册/健康检查外）需携带 Authorization: Bearer <token> 请求头。\n'
               '响应格式统一为：{"code": "0", "message": "ok", "data": {...}}')

    # 4. UI Prototype
    add_heading(doc, '4. 用户界面原型', level=1)
    add_para(doc, '系统采用前后端分离架构，前端基于Vue 3 + Element Plus实现，主要页面包括：')
    add_para(doc, '1. 登录页面(/login)：用户名+密码表单，登录/注册切换\n'
               '2. 邮件主页(/)：左侧文件夹导航 + 中间邮件列表 + 右侧邮件预览，顶部账号切换和操作菜单\n'
               '3. 邮件编辑弹窗：收件人/抄送/密送输入、主题、富文本正文、附件上传\n'
               '4. 设置页面：邮箱账号管理、个人信息编辑')

    path = os.path.join(DOCS_DIR, '2.软件需求规约.docx')
    doc.save(path)
    print(f"  ✓ 软件需求规约 → {path}")


# ============================================================
# 3. Software Architecture Document
# ============================================================
def generate_architecture():
    print("\n📝 生成软件架构文档...")
    doc = Document()

    title = doc.add_heading('Email System 邮件管理系统', level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_para(doc, '软件架构文档', size=14)
    add_para(doc, '版本 1.0 — 2026-07-07', size=10)
    doc.add_page_break()

    add_heading(doc, '修订历史记录', level=1)
    add_table(doc, ['日期', '版本', '说明', '作者'],
              [['2026/07/07', '1.0', '初始架构文档', 'x12w']])

    add_heading(doc, '1. 简介', level=1)
    add_heading(doc, '1.1 目的', level=2)
    add_para(doc, '本文档从架构角度全面描述Email System邮件管理系统的设计，采用"4+1"视图模型，涵盖用例视图、逻辑视图、进程视图、部署视图和实施视图，为开发团队提供统一的架构指导。')

    add_heading(doc, '1.2 范围', level=2)
    add_para(doc, '本文档覆盖Email System的全栈架构设计，包括前端SPA应用、后端Spring Boot服务、数据库层、缓存层、对象存储层、邮件服务层，以及各组件之间的交互关系。')

    add_heading(doc, '1.3 定义与缩略语', level=2)
    add_table(doc, ['术语', '说明'],
              [['SPA', 'Single Page Application，单页应用'],
               ['DTO', 'Data Transfer Object，数据传输对象'],
               ['ORM', 'Object-Relational Mapping，对象关系映射'],
               ['IoC', 'Inversion of Control，控制反转/依赖注入'],
               ['AOP', 'Aspect-Oriented Programming，面向切面编程']])

    add_heading(doc, '2. 架构表示方式', level=1)
    add_para(doc, '系统采用"4+1"视图模型描述架构：\n'
               '• 用例视图：描述系统与外部参与者（用户、外部服务）的交互\n'
               '• 逻辑视图：描述系统的分层结构和核心设计模式\n'
               '• 进程视图：描述系统的并发、定时任务和线程模型\n'
               '• 部署视图：描述系统的物理部署拓扑和容器编排\n'
               '• 实施视图：描述代码组织和模块划分')

    add_heading(doc, '3. 架构目标和约束', level=1)
    add_table(doc, ['目标/约束', '说明'],
              [['前后端分离', '前端Vue 3 SPA + 后端Spring Boot REST API，通过HTTP/JSON通信'],
               ['无状态认证', 'JWT Token实现无状态认证，支持水平扩展'],
               ['数据持久化', 'MySQL作为主数据库，MyBatis-Plus作为ORM框架'],
               ['对象存储', 'MinIO提供S3兼容的对象存储，支持附件管理'],
               ['缓存层', 'Redis用于Token黑名单、会话缓存、速率限制'],
               ['插件架构', '智能分析模块采用插件化设计，支持Java规则引擎和Python插件'],
               ['容器化部署', 'Docker Compose编排全部服务，支持一键部署']])

    add_heading(doc, '4. 用例视图', level=1)
    add_para(doc, '核心用例实现流程：\n\n'
               '用户登录流程：前端LoginView → POST /api/auth/login → AuthController → UserRegistryService → '
               'BCrypt密码校验 → TokenService签发JWT → 返回{accessToken, refreshToken} → 前端存储并跳转主页\n\n'
               'IMAP同步流程：@Scheduled定时触发 → MailSyncService.syncAllAccounts() → 遍历所有用户 → '
               '线程池提交同步任务 → Jakarta Mail连接IMAP服务器 → 拉取INBOX邮件 → Message-ID去重 → '
               '解析邮件内容(From/Subject/Body/Attachments) → 持久化到MySQL → 更新同步时间戳\n\n'
               '邮件发送流程：前端编辑邮件 → POST /api/messages/send → MailMessageController → '
               'MailSenderService.send() → JavaMailSender通过SMTP发送 → 持久化到已发送文件夹 → 返回邮件详情')

    add_heading(doc, '5. 逻辑视图', level=1)
    add_heading(doc, '5.1 分层架构', level=2)
    add_para(doc, '后端采用经典三层架构 + 横切关注点：\n\n'
               '┌──────────────────────────────────────┐\n'
               '│   Controller 层 (REST API)           │\n'
               '│   AuthController, MailMessageController,│\n'
               '│   MailAccountController, ...          │\n'
               '├──────────────────────────────────────┤\n'
               '│   Service 层 (业务逻辑)              │\n'
               '│   MailSenderService, MailSyncService,  │\n'
               '│   UserRegistryService, ...             │\n'
               '├──────────────────────────────────────┤\n'
               '│   Mapper 层 (数据访问)               │\n'
               '│   MyBatis-Plus BaseMapper             │\n'
               '├──────────────────────────────────────┤\n'
               '│   Entity 层 (数据模型)               │\n'
               '│   MailMessage, MailAccount, SysUser... │\n'
               '└──────────────────────────────────────┘\n'
               '横切关注点：Security（JWT Filter）、Exception Handling（GlobalExceptionHandler）、'
               'Rate Limiting（RateLimitFilter）、Validation（Jakarta Validation）')

    add_heading(doc, '5.2 核心设计包', level=2)
    add_table(doc, ['包路径', '职责'],
              [['com.example.emailsystem.controller', 'REST API控制器，接收HTTP请求并返回JSON响应'],
               ['com.example.emailsystem.service', '业务逻辑层，封装核心业务规则和流程'],
               ['com.example.emailsystem.mapper', 'MyBatis-Plus数据访问接口'],
               ['com.example.emailsystem.entity', 'JPA实体类，映射数据库表结构'],
               ['com.example.emailsystem.dto', '数据传输对象，API请求/响应模型'],
               ['com.example.emailsystem.security', '安全组件：JWT、加密、认证过滤器'],
               ['com.example.emailsystem.config', 'Spring配置：Security、CORS、Rate Limit'],
               ['com.example.emailsystem.common', '通用组件：API响应封装、异常处理、校验'],
               ['com.example.emailsystem.intelligence', '智能分析插件系统']])

    add_heading(doc, '6. 进程视图', level=1)
    add_para(doc, '系统进程/线程模型：\n\n'
               '• HTTP请求处理线程：Tomcat默认线程池(200线程)，处理REST API请求\n'
               '• IMAP同步线程池：固定4线程(ExecutorService)，每60秒触发一次同步任务\n'
               '• 定时任务线程：Spring @Scheduled线程池，执行周期性任务\n'
               '   - syncAllAccounts(): 每60秒，遍历所有用户触发IMAP同步\n'
               '   - evictExpiredEntries(): 每5分钟，清理速率限制过期条目\n'
               '• Redis连接池：Lettuce连接池管理Redis连接\n'
               '• 数据库连接池：HikariCP连接池管理MySQL连接\n\n'
               '并发安全：\n'
               '• RateLimitFilter使用ConcurrentHashMap保证线程安全\n'
               '• IntelligenceAnalysisServiceImpl使用ConcurrentHashMap缓存分析结果\n'
               '• MailSyncService使用线程池隔离同步任务')

    add_heading(doc, '7. 部署视图', level=1)
    add_para(doc, '生产环境部署拓扑（Docker Compose）：\n\n'
               '┌─────────────────────────────────────────────┐\n'
               '│  Nginx (:80) — 静态文件服务 + API反向代理   │\n'
               '├─────────────────────────────────────────────┤\n'
               '│  Backend (:8080) — Spring Boot应用          │\n'
               '├────────────────┬──────────────┬─────────────┤\n'
               '│  MySQL (:3306)  │ Redis (:6379) │ MinIO (:9000)│\n'
               '├────────────────┴──────────────┴─────────────┤\n'
               '│  Mailpit (:1025) — 邮件测试                 │\n'
               '│  Mailserver (:25/:993) — 独立邮件服务       │\n'
               '└─────────────────────────────────────────────┘\n\n'
               '数据持久化：\n'
               '• MySQL数据卷: ./data/mysql → /var/lib/mysql\n'
               '• Redis数据卷: ./data/redis → /data\n'
               '• MinIO数据卷: ./data/uploads → /data\n'
               '• SSL证书只读挂载: /etc/letsencrypt → 容器内')

    add_heading(doc, '8. 实施视图', level=1)
    add_para(doc, '项目代码组织结构：\n\n'
               'email-system/\n'
               '├── backend/           # Spring Boot后端\n'
               '│   ├── src/main/java/com/example/emailsystem/\n'
               '│   ├── src/main/resources/application.yml\n'
               '│   └── pom.xml\n'
               '├── frontend/          # Vue 3前端\n'
               '│   ├── src/views/     # 页面组件\n'
               '│   ├── src/api/       # API客户端\n'
               '│   ├── src/stores/    # Pinia状态管理\n'
               '│   └── src/router/    # Vue Router路由\n'
               '├── plugins/           # 智能分析插件\n'
               '│   └── intelligence/python/\n'
               '├── deploy/nginx/      # Nginx配置\n'
               '├── scripts/           # 部署与管理脚本\n'
               '├── docker-compose.yml # 容器编排\n'
               '└── .env.example       # 环境变量模板')

    add_heading(doc, '9. 数据视图', level=1)
    add_para(doc, '核心数据表设计：\n\n'
               '• sys_user — 用户表(id, username, password_hash, display_name, email, status)\n'
               '• mail_account — 邮箱账号表(id, user_id, email_address, smtp_host, smtp_port, imap_host, imap_port, auth_username, auth_password_encrypted)\n'
               '• mail_message — 邮件表(id, user_id, account_id, folder_id, message_id, from_address, subject, content, preview, read_flag, deleted_flag)\n'
               '• mail_folder — 文件夹表(id, account_id, name, type, unread_count, total_count)\n'
               '• mail_recipient — 收件人表(id, message_id, type[to/cc/bcc], email_address, display_name)\n'
               '• mail_attachment — 附件表(id, message_id, original_name, content_type, size_bytes, storage_type, storage_path)\n'
               '• mail_contact — 联系人表(id, user_id, name, email_address, phone, remark)\n'
               '• mail_push_event — 推送事件表(id, user_id, message_id, event_type, title, content, priority, read_flag)\n'
               '• mail_intelligence_result — 智能分析结果表(id, message_id, spam_label, priority_label, risk_level, plugin_name)')

    add_heading(doc, '10. 质量属性', level=1)
    add_table(doc, ['质量属性', '实现策略'],
              [['安全性', 'JWT无状态认证 + BCrypt密码哈希 + AES-256-GCM邮箱密码加密 + Redis Token黑名单 + IP速率限制 + CORS白名单'],
               ['可靠性', '数据库事务保证数据一致性 + 全局异常处理 + 健康检查端点'],
               ['可维护性', '分层架构 + 插件化设计 + 统一API响应格式 + 代码注释'],
               ['可部署性', 'Docker Compose一键部署 + 自动化健康检查 + 环境变量配置化'],
               ['可测试性', '前后端分离独立测试 + Spring Boot Test支持 + REST API可测试性']])

    path = os.path.join(DOCS_DIR, '5.软件架构文档.docx')
    doc.save(path)
    print(f"  ✓ 软件架构文档 → {path}")


# ============================================================
# 4. Use Case Specification
# ============================================================
def generate_use_cases():
    print("\n📝 生成用例规约...")
    doc = Document()

    title = doc.add_heading('Email System — 用例规约', level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_para(doc, '版本 1.0 — 2026-07-07', size=10)
    doc.add_page_break()

    # Use Case 1: User Login
    add_heading(doc, '用例1：用户登录', level=1)
    add_heading(doc, '1.1 简要说明', level=2)
    add_para(doc, '注册用户通过用户名和密码登录系统，获取JWT Token以访问受保护资源。')

    add_heading(doc, '1.2 事件流', level=2)
    add_heading(doc, '1.2.1 基本流', level=3)
    add_para(doc, '1. 用户访问登录页面(/login)\n'
               '2. 用户输入用户名和密码，点击登录按钮\n'
               '3. 前端发送POST /api/auth/login请求，携带{username, password}\n'
               '4. 后端UserRegistryService验证用户名和密码(BCrypt)\n'
               '5. 验证通过后TokenService签发AccessToken和RefreshToken\n'
               '6. 返回{accessToken, refreshToken, expiresIn, user}\n'
               '7. 前端存储Token到localStorage，跳转到邮件主页(/)')

    add_heading(doc, '1.2.2 备选流', level=3)
    add_para(doc, 'A1. 用户名或密码错误：返回AUTH_401错误码，前端显示"用户名或密码错误"\n'
               'A2. 速率限制触发：连续5次失败后返回429 Too Many Requests，60秒后自动解除\n'
               'A3. Token过期：前端拦截器检测401响应，自动跳转登录页')

    add_heading(doc, '1.3 前置条件', level=2)
    add_para(doc, '• 用户已注册账号\n• 后端服务正常运行\n• Redis服务可用(Token黑名单检查)')

    add_heading(doc, '1.4 后置条件', level=2)
    add_para(doc, '• 用户获得有效的JWT Token\n• 后续API请求携带Authorization: Bearer <token>')

    doc.add_page_break()

    # Use Case 2: IMAP Sync
    add_heading(doc, '用例2：IMAP邮件同步', level=1)
    add_heading(doc, '2.1 简要说明', level=2)
    add_para(doc, '系统定时自动从配置的IMAP服务器拉取邮件，按Message-ID去重后存入本地数据库。')

    add_heading(doc, '2.2 事件流', level=2)
    add_heading(doc, '2.2.1 基本流', level=3)
    add_para(doc, '1. Spring @Scheduled定时器每60秒触发syncAllAccounts()\n'
               '2. 遍历所有活跃用户及其邮箱账号\n'
               '3. 线程池提交同步任务(最多4并发)\n'
               '4. 连接IMAP服务器(Jakarta Mail IMAPS协议)\n'
               '5. 打开INBOX文件夹(READ_WRITE模式)\n'
               '6. 首次同步拉取最近50封(可配置)，后续增量拉取20封\n'
               '7. 遍历邮件：解析Message-ID → 本地去重 → 提取From/Subject/Body → 保存到MySQL\n'
               '8. 标记已同步邮件为SEEN\n'
               '9. 更新账号lastSyncAt时间戳\n'
               '10. 更新文件夹邮件计数')

    add_heading(doc, '2.2.2 备选流', level=3)
    add_para(doc, 'A1. IMAP连接失败：记录错误日志，跳过该账号，不影响其他账号同步\n'
               'A2. From地址为空：安全处理，使用空字符串作为发件人地址\n'
               'A3. HTML内容解析异常：返回空字符串作为邮件正文')

    add_heading(doc, '2.3 前置条件', level=2)
    add_para(doc, '• 邮箱账号已配置IMAP服务器信息\n• IMAP服务器可访问\n• 账号密码已解密')

    add_heading(doc, '2.4 后置条件', level=2)
    add_para(doc, '• 新邮件已持久化到mail_message表\n• 文件夹计数已更新\n• 同步时间戳已更新')

    doc.add_page_break()

    # Use Case 3: Send Mail
    add_heading(doc, '用例3：发送邮件', level=1)
    add_heading(doc, '3.1 简要说明', level=2)
    add_para(doc, '用户编写邮件并通过SMTP协议发送给指定收件人，同时保存副本到已发送文件夹。')

    add_heading(doc, '3.2 事件流', level=2)
    add_heading(doc, '3.2.1 基本流', level=3)
    add_para(doc, '1. 用户点击"写邮件"按钮，打开发送邮件表单\n'
               '2. 填写收件人(To)/抄送(CC)/密送(BCC)、主题、正文(HTML)\n'
               '3. 可选：上传附件\n'
               '4. 点击"发送"按钮\n'
               '5. 前端POST /api/messages/send，携带SendMessageRequest\n'
               '6. 后端MailSenderService:\n'
               '   a. 构建JavaMailSender(SMTP配置)\n'
               '   b. 创建MimeMessage，设置From/To/CC/BCC/Subject/Content/Attachments\n'
               '   c. 通过SMTP发送邮件\n'
               '   d. 持久化邮件到已发送文件夹\n'
               '   e. 保存收件人信息\n'
               '7. 返回发送成功的邮件详情')

    add_heading(doc, '3.2.2 备选流', level=3)
    add_para(doc, 'A1. SMTP发送失败：抛出BusinessException，前端显示"邮件发送失败"\n'
               'A2. 保存草稿：点击"存草稿"而非"发送"，邮件保存到草稿文件夹(draftFlag=1)\n'
               'A3. 附件过大：超过10MB限制，返回错误提示')

    add_heading(doc, '3.3 前置条件', level=2)
    add_para(doc, '• 用户已登录\n• 已配置至少一个邮箱账号(SMTP信息完整)\n• SMTP服务器可访问')

    add_heading(doc, '3.4 后置条件', level=2)
    add_para(doc, '• 邮件已通过SMTP发送给收件人\n• 邮件副本已保存到已发送文件夹')

    doc.add_page_break()

    # Use Case 4: Intelligent Analysis
    add_heading(doc, '用例4：智能邮件分析', level=1)
    add_heading(doc, '4.1 简要说明', level=2)
    add_para(doc, '对指定邮件进行智能分析，检测垃圾邮件、分析优先级、识别安全威胁。')

    add_heading(doc, '4.2 事件流', level=2)
    add_heading(doc, '4.2.1 基本流', level=3)
    add_para(doc, '1. 用户查看邮件详情\n'
               '2. 前端GET /api/intelligence/messages/{id}/analyze\n'
               '3. 后端IntelligenceAnalysisServiceImpl:\n'
               '   a. 获取邮件内容(subject + plainText + html + links)\n'
               '   b. 调用插件client.analyzeEmailJson(requestJson)\n'
               '   c. RuleBasedIntelligencePluginClient执行规则分析:\n'
               '      - 关键词匹配(urgent/紧急/审批/中奖/免费/casino...)\n'
               '      - 链接检测(IP地址URL、敏感关键词login/password/pay)\n'
               '      - 评分计算(priority/spam/risk)\n'
               '   d. 解析插件返回的JSON结果\n'
               '   e. 缓存分析结果\n'
               '4. 返回IntelligenceAnalysisResult(含spam/priority/risk/threats)')

    add_heading(doc, '4.2.2 备选流', level=3)
    add_para(doc, 'A1. 插件不可用：返回默认结果(spam=unknown, priority=normal, risk=none)\n'
               'A2. 缓存命中：直接返回已缓存的分析结果，无需重新分析\n'
               'A3. 缓存超限(1000条)：淘汰最旧条目后存储新结果')

    add_heading(doc, '4.3 特殊需求', level=2)
    add_para(doc, '• 插件执行超时时间：2秒(可配置)\n'
               '• 分析结果缓存上限：1000条\n'
               '• 高风险邮件自动生成推送安全告警事件')

    path = os.path.join(DOCS_DIR, '6.用例规约模板.docx')
    doc.save(path)
    print(f"  ✓ 用例规约 → {path}")


# ============================================================
# 5. Project Development Summary Report
# ============================================================
def generate_summary():
    print("\n📝 生成项目开发总结报告...")
    doc = Document()

    title = doc.add_heading('Email System 邮件管理系统', level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_para(doc, '项目开发总结报告', size=14)
    add_para(doc, '版本 1.0 — 2026-07-07', size=10)
    doc.add_page_break()

    add_heading(doc, '1. 项目概述', level=1)
    add_para(doc, 'Email System是一个基于Spring Boot + Vue 3的全栈邮件管理平台，支持多邮箱账号管理、IMAP邮件同步、'
               'SMTP邮件发送、附件管理、联系人管理以及智能邮件分析等功能。项目采用前后端分离架构，通过Docker Compose实现'
               '一键部署。')

    add_heading(doc, '2. 项目进度', level=1)
    add_table(doc, ['阶段', '时间', '主要工作'],
              [['需求分析与设计', '2026.05.17 - 2026.05.26', '数据库设计、API规范定义、技术选型'],
               ['MVP开发', '2026.05.26 - 2026.06.08', '前端框架搭建、后端三层架构、认证模块、邮件收发'],
               ['功能完善', '2026.06.09 - 2026.06.22', 'IMAP同步、附件管理、联系人、前端交互完善'],
               ['智能化与部署', '2026.06.23 - 2026.06.29', '智能分析插件、部署脚本、Docker编排'],
               ['安全加固与文档', '2026.06.30 - 2026.07.07', '安全漏洞修复、文档编写、Bug修复']])

    add_heading(doc, '3. 技术栈', level=1)
    add_table(doc, ['层次', '技术选型', '版本'],
              [['后端框架', 'Spring Boot', '3.4.1'],
               ['ORM', 'MyBatis-Plus', '3.5.9'],
               ['数据库', 'MySQL', '8.4'],
               ['缓存', 'Redis', '7.4'],
               ['对象存储', 'MinIO', 'RELEASE.2025-02-28'],
               ['前端框架', 'Vue 3 + Vite', '3.x + 5.x'],
               ['UI库', 'Element Plus', '最新'],
               ['认证', 'JWT (HMAC-SHA256)', '自定义实现'],
               ['邮件协议', 'Jakarta Mail (SMTP/IMAP)', '—'],
               ['容器化', 'Docker + Docker Compose', '—'],
               ['Web服务器', 'Nginx', '1.27-alpine']])

    add_heading(doc, '4. 核心功能清单', level=1)
    add_table(doc, ['模块', '功能点', '状态'],
              [['用户认证', '注册/登录/Token刷新/安全登出', '✅ 完成'],
               ['邮箱管理', '多账号CRUD/SMTP-IMAP配置/账号切换', '✅ 完成'],
               ['邮件收发', 'SMTP发送(To/CC/BCC/HTML/附件)/IMAP同步接收', '✅ 完成'],
               ['文件夹', '收件箱/发件箱/草稿箱/已删除/邮件计数', '✅ 完成'],
               ['联系人', '通讯录增删改查', '✅ 完成'],
               ['附件', '上传/下载/删除(MinIO+本地存储)', '✅ 完成'],
               ['智能分析', '垃圾检测/优先级/威胁识别/推送事件', '✅ 完成'],
               ['安全', 'JWT/BCrypt/AES-256-GCM/速率限制/Token黑名单', '✅ 完成'],
               ['部署', 'Docker Compose/deploy.sh一键部署/健康检查', '✅ 完成'],
               ['文档', '需求规约/架构文档/用例规约/周报/PPT', '✅ 完成']])

    add_heading(doc, '5. 关键指标', level=1)
    add_table(doc, ['指标', '数值'],
              [['后端Java文件数', '50+'],
               ['前端Vue/TS文件数', '15+'],
               ['数据库表', '10'],
               ['REST API端点', '30+'],
               ['Docker服务', '7 (MySQL/Redis/MinIO/Mailpit/Mailserver/Backend/Nginx)'],
               ['开发周期', '约51天 (2026.05.17 - 2026.07.07)'],
               ['Git提交数', '50+'],
               ['代码行数(估算)', '~8000+']])

    add_heading(doc, '6. 经验与教训', level=1)
    add_para(doc, '成功经验：\n'
               '• 前后端分离架构使开发和调试更加独立高效\n'
               '• Docker Compose编排大幅简化了环境搭建和部署流程\n'
               '• 插件化智能分析设计便于后续扩展新的分析能力\n'
               '• 尽早编写部署脚本(deploy.sh)避免了手工操作的繁琐\n\n'
               '待改进：\n'
               '• 单元测试覆盖率不足，后续应加强自动化测试\n'
               '• 前端状态管理较为简单，复杂场景可考虑更完善的状态管理方案\n'
               '• API接口缺少Swagger文档，前后端协作效率可进一步提升\n'
               '• 日志收集和监控体系尚未建立，生产环境可观测性不足')

    add_heading(doc, '7. 后续规划', level=1)
    add_para(doc, '• 邮件全文搜索（Elasticsearch集成）\n'
               '• 邮件规则引擎（自动归档、标签分类、自动回复）\n'
               '• 多用户管理后台和RBAC权限控制\n'
               '• API文档（Swagger/OpenAPI）\n'
               '• 单元测试和集成测试补充\n'
               '• CI/CD流水线搭建\n'
               '• 监控和告警（Prometheus + Grafana）\n'
               '• 移动端适配或PWA支持')

    path = os.path.join(DOCS_DIR, '7.项目开发总结报告.docx')
    doc.save(path)
    print(f"  ✓ 项目开发总结报告 → {path}")


# ============================================================
# Main
# ============================================================
if __name__ == '__main__':
    print("=" * 60)
    print("Email System 项目文档生成器")
    print("=" * 60)

    generate_all_weekly_reports()
    generate_srs()
    generate_architecture()
    generate_use_cases()
    generate_summary()

    print("\n" + "=" * 60)
    print("✅ 全部文档生成完毕!")
    print("=" * 60)
